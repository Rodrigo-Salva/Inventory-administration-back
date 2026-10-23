"""Generación de reportes Excel/CSV con el estilo corporativo común."""

import csv
import io
from datetime import datetime
from typing import Iterable, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

DEFAULT_TENANT_TITLE = "INVENTORY ADMINISTRATION"
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
_BORDER_SIDE = Side(style="thin", color="DDDDDD")
_BORDER = Border(left=_BORDER_SIDE, right=_BORDER_SIDE, top=_BORDER_SIDE, bottom=_BORDER_SIDE)


def safe_cell(value):
    """Neutraliza inyección de fórmulas (=, +, -, @) en texto controlado por usuarios."""
    if isinstance(value, str) and value.startswith(_FORMULA_PREFIXES):
        return "'" + value
    return value


def _range_label(start: datetime | None, end: datetime | None, empty_label: str) -> str:
    if start and end:
        return f"Rango: {start.strftime('%d/%m/%Y')} al {end.strftime('%d/%m/%Y')}"
    return empty_label


def build_excel(
    *,
    sheet: str,
    title: str,
    tenant_name: str | None,
    headers: Sequence[str],
    rows: Iterable[Sequence],
    width: int = 20,
    column_widths: dict[str, int] | None = None,
    money_columns: Sequence[int] = (),
    show_range: bool = False,
    start: datetime | None = None,
    end: datetime | None = None,
    empty_range_label: str = "Todos los registros",
) -> io.BytesIO:
    """Construye el libro: título, subtítulo opcional (rango), cabecera y filas."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    last_col = get_column_letter(len(headers))

    ws.merge_cells(f"A1:{last_col}1")
    ws["A1"] = f"{title} - {tenant_name.upper() if tenant_name else DEFAULT_TENANT_TITLE}"
    ws["A1"].font = Font(name="Arial", size=16, bold=True, color="1E40AF")
    ws["A1"].alignment = Alignment(horizontal="center")

    if show_range:
        ws.merge_cells(f"A2:{last_col}2")
        ws["A2"] = (
            f"Generado el: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')} | "
            f"{_range_label(start, end, empty_range_label)}"
        )
        ws["A2"].font = Font(size=10, italic=True)
        ws["A2"].alignment = Alignment(horizontal="center")

    header_row = 4 if show_range else 3
    ws.append([])
    ws.append(list(headers))
    for cell in ws[header_row]:
        cell.font = Font(name="Arial", size=12, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _BORDER

    for index, row in enumerate(rows, start=header_row + 1):
        ws.append([safe_cell(v) for v in row])
        for cell in ws[index]:
            cell.border = _BORDER
            if cell.column in money_columns:
                cell.number_format = '"$"#,##0.00'

    for column in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(column)].width = width
    for letter, value in (column_widths or {}).items():
        ws.column_dimensions[letter].width = value

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def build_csv(headers: Sequence[str], rows: Iterable[Sequence]) -> str:
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(headers)
    for row in rows:
        writer.writerow([safe_cell(v) for v in row])
    return output.getvalue()
